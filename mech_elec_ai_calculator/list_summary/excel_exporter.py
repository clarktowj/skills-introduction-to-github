import os
from typing import List
from common.models import ListSummary, QuantityItem
from framework.exceptions import FileError

class ExcelExporter:
    def __init__(self):
        pass
    
    def export(self, summary: ListSummary, output_path: str) -> str:
        try:
            try:
                import openpyxl
                from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
                from openpyxl.utils import get_column_letter
            except ImportError:
                return self._export_csv(summary, output_path)
            
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "工程量清单"
            
            self._write_header(ws)
            self._write_items(ws, summary.items)
            self._write_summary(ws, summary)
            
            directory = os.path.dirname(output_path)
            if directory and not os.path.exists(directory):
                os.makedirs(directory, exist_ok=True)
            
            wb.save(output_path)
            return output_path
        
        except Exception as e:
            raise FileError(f"Failed to export Excel file: {str(e)}")
    
    def _export_csv(self, summary: ListSummary, output_path: str) -> str:
        try:
            directory = os.path.dirname(output_path)
            if directory and not os.path.exists(directory):
                os.makedirs(directory, exist_ok=True)
            
            csv_path = output_path.replace('.xlsx', '.csv')
            
            with open(csv_path, 'w', encoding='utf-8-sig') as f:
                f.write("序号,编号,项目名称,规格型号,单位,数量,单价,合价,分类,子分类\n")
                
                for i, item in enumerate(summary.items, 1):
                    f.write(f"{i},{item.code},{item.name},,{item.unit},{item.quantity},{item.unit_price},{item.total_price},{item.category},{item.sub_category}\n")
                
                f.write(f",,,,,,,{summary.total_cost},合计,\n")
            
            return csv_path
        except Exception as e:
            raise FileError(f"Failed to export CSV file: {str(e)}")
    
    def _write_header(self, ws) -> None:
        headers = ['序号', '定额编号', '项目名称', '规格型号', '计量单位', '工程量', '单价(元)', '合价(元)', '分类', '子分类', '备注']
        
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        header_alignment = Alignment(horizontal='center', vertical='center')
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
        
        ws.freeze_panes = 'A2'
    
    def _write_items(self, ws, items: List[QuantityItem]) -> None:
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        alignment = Alignment(horizontal='center', vertical='center')
        
        for row, item in enumerate(items, 2):
            ws.cell(row=row, column=1, value=row-1).alignment = alignment
            ws.cell(row=row, column=2, value=item.code).alignment = alignment
            ws.cell(row=row, column=3, value=item.name)
            ws.cell(row=row, column=4, value=item.attributes.get('spec', ''))
            ws.cell(row=row, column=5, value=item.unit).alignment = alignment
            ws.cell(row=row, column=6, value=item.quantity).alignment = alignment
            ws.cell(row=row, column=7, value=item.unit_price).alignment = alignment
            ws.cell(row=row, column=8, value=item.total_price).alignment = alignment
            ws.cell(row=row, column=9, value=item.category).alignment = alignment
            ws.cell(row=row, column=10, value=item.sub_category).alignment = alignment
            ws.cell(row=row, column=11, value=item.description)
            
            for col in range(1, 12):
                ws.cell(row=row, column=col).border = thin_border
        
        for col in range(1, 12):
            ws.column_dimensions[get_column_letter(col)].auto_size = True
    
    def _write_summary(self, ws, summary: ListSummary) -> None:
        last_row = len(summary.items) + 2
        
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        ws.cell(row=last_row, column=1, value='合计').border = thin_border
        ws.cell(row=last_row, column=8, value=summary.total_cost).border = thin_border
        
        ws.cell(row=last_row+1, column=1, value='项目名称').border = thin_border
        ws.cell(row=last_row+1, column=2, value=summary.project_name).border = thin_border
        
        ws.cell(row=last_row+2, column=1, value='项目编号').border = thin_border
        ws.cell(row=last_row+2, column=2, value=summary.project_code).border = thin_border